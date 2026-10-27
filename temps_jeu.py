"""
Page « Temps de jeu » à partir de l'export de l'appli temps-de-jeu-d1 (TJ_<...>.xlsx).

Onglets lus : « Temps de jeu » (Numéro, Nom, Gardien, Temps 1re MT, Temps 2e MT, Entrées)
              « Match » (Équipe, Adversaire)
Le total est toujours recalculé (1re MT + 2e MT) : la colonne « Total (s) » du fichier n'est pas utilisée.
L'onglet « Entrées-Sorties » (journal de l'appli) servira plus tard aux cinq sur le terrain à chaque instant.
"""
import re
import openpyxl
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches
import matplotlib.lines

import cartes as C

ALIAS_EQUIPE = {"EDF": "FRANCE"}


class TempsJeuError(Exception):
    """Fichier de temps de jeu illisible."""


def _sec(v):
    """'08:30' -> 510 ; nombre de secondes accepté aussi."""
    if v is None or str(v).strip() == "":
        return 0
    if isinstance(v, (int, float)):
        return int(v)
    m = re.fullmatch(r"\s*(\d+):(\d{1,2})\s*", str(v))
    if not m:
        raise TempsJeuError(f"temps illisible : « {v} » (format attendu mm:ss)")
    return int(m.group(1)) * 60 + int(m.group(2))


def mmss(s):
    return f"{int(s) // 60:02d}:{int(s) % 60:02d}"


def lire(path):
    """Renvoie dict(equipe, adversaire, joueurs=[...], alertes=[...])."""
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        raise TempsJeuError(f"fichier illisible : {e}")
    if "Temps de jeu" not in wb.sheetnames:
        raise TempsJeuError("onglet « Temps de jeu » introuvable : utilise l'export de l'appli temps de jeu")
    info = {}
    if "Match" in wb.sheetnames:
        for r in wb["Match"].iter_rows(values_only=True):
            if r and r[0]:
                info[str(r[0]).strip()] = r[1]
    ws = wb["Temps de jeu"]
    rows = list(ws.iter_rows(values_only=True))
    tete = [str(h).strip() if h else "" for h in rows[0]]
    col = {h: i for i, h in enumerate(tete)}
    for h in ("Numéro", "Nom", "Temps 1re MT", "Temps 2e MT"):
        if h not in col:
            raise TempsJeuError(f"colonne « {h} » absente de l'onglet Temps de jeu")
    joueurs, alertes = [], []
    for r in rows[1:]:
        if not r or r[col["Nom"]] is None:
            continue
        j = dict(num=str(r[col["Numéro"]]).strip(), nom=str(r[col["Nom"]]).strip().upper(),
                 gardien=str(r[col["Gardien"]]).strip().upper() == "OUI" if "Gardien" in col else False,
                 mt1=_sec(r[col["Temps 1re MT"]]), mt2=_sec(r[col["Temps 2e MT"]]),
                 entrees=int(r[col["Entrées"]] or 0) if "Entrées" in col else 0)
        j["total"] = j["mt1"] + j["mt2"]
        if "Total (s)" in col and r[col["Total (s)"]] not in (None, "") and int(r[col["Total (s)"]]) != j["total"]:
            alertes.append(f"Temps de jeu {j['nom']} : « Total (s) » = {r[col['Total (s)']]} dans le fichier, "
                           f"1re + 2e MT = {j['total']} (c'est cette valeur qui est utilisée)")
        if j["mt1"] > 1200 or j["mt2"] > 1200:
            alertes.append(f"Temps de jeu {j['nom']} : plus de 20:00 sur une mi-temps")
        joueurs.append(j)
    # vraisemblance par mi-temps : 20:00 de gardien, environ 4 × 20:00 de joueurs de champ
    for k, mt in ((1, "mt1"), (2, "mt2")):
        g = sum(j[mt] for j in joueurs if j["gardien"])
        c = sum(j[mt] for j in joueurs if not j["gardien"])
        if g and abs(g - 1200) > 30:
            alertes.append(f"Temps de jeu {k}re MT : gardiens {mmss(g)} au total (20:00 attendu)".replace("2re", "2e"))
        if c and abs(c - 4800) > 120:
            alertes.append(f"Temps de jeu {k}re MT : joueurs de champ {mmss(c)} au total (environ 80:00 attendu, "
                           f"hors power play)".replace("2re", "2e"))
    eq = str(info.get("Équipe") or "").strip().upper()
    return dict(equipe=ALIAS_EQUIPE.get(eq, eq), adversaire=str(info.get("Adversaire") or "").strip().upper(),
                joueurs=joueurs, alertes=alertes)


def page(pdf, num, tj, couleur):
    """Une ligne par joueur : barre 1re MT + 2e MT, puis 1re MT, 2e MT, total, entrées, durée moyenne d'une entrée."""
    fig = plt.figure(figsize=(7.874, 7.874), dpi=200)
    C.cadre_page(fig, "Temps de jeu", "Qui a joué, et combien", num)
    champ = sorted([j for j in tj["joueurs"] if not j["gardien"] and j["total"] > 0], key=lambda j: -j["total"])
    gks = sorted([j for j in tj["joueurs"] if j["gardien"] and j["total"] > 0], key=lambda j: -j["total"])
    lignes = champ + gks
    c2 = tuple(0.45 * np.array(matplotlib.colors.to_rgb(couleur)) + 0.55)   # même couleur, plus claire, opaque

    # colonnes (abscisses en fraction de page, alignement à droite)
    COLS = [("1re MT", 0.625), ("2e MT", 0.70), ("TOTAL", 0.785), ("ENTRÉES", 0.865), ("MOY.", 0.935)]
    yh = 0.80
    fig.text(0.07, yh, tj["equipe"], fontproperties=C.PLAY, fontsize=15, color=couleur, va="bottom")
    for lab, x in COLS:
        fig.text(x, yh + 0.003, lab, fontproperties=C.POPB, fontsize=7, color=C.GRIS, ha="right", va="bottom")
    fig.add_artist(matplotlib.lines.Line2D([0.07, 0.935], [yh - 0.006, yh - 0.006], transform=fig.transFigure,
                                           c=C.ENCRE, lw=1.2))

    ax = fig.add_axes([0.255, 0.15, 0.295, 0.635])
    n = len(lignes)
    y = np.arange(n)[::-1].astype(float)
    if gks:  # espace avant les gardiens
        y[len(champ):] -= 0.7
    mx = max(j["total"] for j in lignes)
    for k, (yi, j) in enumerate(zip(y, lignes)):
        ax.barh(yi, j["mt1"], color=couleur, height=0.78, zorder=3)
        ax.barh(yi, j["mt2"], left=j["mt1"], color=c2, height=0.78, zorder=3, edgecolor="white", linewidth=0.8)
    ax.set_ylim(min(y) - 0.55, max(y) + 0.55)
    ax.set_xlim(0, mx * 1.03)
    ax.axis("off")

    tr = ax.transData
    demi = (fig.transFigure.inverted().transform(tr.transform((0, 0.5)))[1]
            - fig.transFigure.inverted().transform(tr.transform((0, 0)))[1])
    for k, (yi, j) in enumerate(zip(y, lignes)):
        yf = fig.transFigure.inverted().transform(tr.transform((0, yi)))[1]
        if k % 2 == 0:   # une ligne sur deux teintée, pour suivre la ligne à l'œil
            fig.add_artist(matplotlib.patches.Rectangle((0.065, yf - demi), 0.875, 2 * demi,
                                                        transform=fig.transFigure, fc="#EAE6D9", ec="none",
                                                        zorder=-5))
        fig.text(0.07, yf, j["num"], fontproperties=C.POPB, fontsize=9, color=couleur, va="center")
        fig.text(0.10, yf, j["nom"] + (" (GB)" if j["gardien"] else ""), fontproperties=C.POPS, fontsize=9,
                 color=C.ENCRE, va="center")
        moy = mmss(j["total"] / j["entrees"]) if j["entrees"] else "–"
        vals = [(mmss(j["mt1"]), C.POPS, 9, "#3A3936"), (mmss(j["mt2"]), C.POPS, 9, "#3A3936"),
                (mmss(j["total"]), C.PLAY, 13, C.ENCRE), (str(j["entrees"]), C.POPS, 9, "#3A3936"),
                (moy, C.POPS, 9, "#3A3936")]
        for (txt, fp, fs, col), (_, x) in zip(vals, COLS):
            fig.text(x - (0.025 if txt == str(j["entrees"]) and x == 0.865 else 0), yf, txt, fontproperties=fp,
                     fontsize=fs, color=col, ha="right", va="center")

    # légende
    for x, c, lab in ((0.255, couleur, "1re MT"), (0.345, c2, "2e MT")):
        fig.add_artist(matplotlib.patches.Rectangle((x, 0.108), 0.016, 0.016, transform=fig.transFigure,
                                                    fc=c, ec="none"))
        fig.text(x + 0.022, 0.116, lab, fontproperties=C.POPS, fontsize=8, color=C.ENCRE, va="center")
    fig.text(0.07, 0.083, "Entrées : nombre de fois où le joueur est entré sur le terrain · "
             "Moy. : durée moyenne d'une entrée", fontproperties=C.POP, fontsize=6.5, color=C.GRIS)
    pdf.savefig(fig, facecolor=C.CREME)
    plt.close(fig)
