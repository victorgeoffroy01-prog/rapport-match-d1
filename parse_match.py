"""
Parsing + calcul des stats de match D1 Futsal à partir de la feuille brute (export outil de tag)
et d'une compo (xlsx, 1 feuille par équipe : Nom / Gardien / Numéro).

Principe général observé sur la feuille brute :
- Chaque ligne = un type d'événement, nommé "<Catégorie> <ÉQUIPE>".
- Certaines catégories sont "bilatérales" : la ligne tague à la fois le joueur auteur (de l'équipe nommée)
  ET son vis-à-vis (équipe adverse), avec un sens différent pour chacun :
    Duel Gagné/Perdu OFF, Faute Subie, Interception, Récupération, Tir Cadré, But.
- D'autres catégories sont "solo" (un seul camp tagué) : Perte de balle, Ballon Rendu, Passe Loupée,
  Passe Décisive, Poteau, Relances gardien, Tir HC, Tir Contré.
- D'autres enfin sont des tags de volume au niveau équipe uniquement (pas de joueur) : ATT 4-0/1-2-1/générique,
  Trans OFF, Touche OFF/DEF, Corner, Coup Franc, Sortie de Jeu, Gardien Volant, Power Play.
"""
import openpyxl
import re
from collections import defaultdict

RAW_PATH = "LAVAL_-_NICE.xlsx"
COMPO_PATH = "Compo_D1_.xlsx"


def load_compo(path):
    """Colonnes attendues : Nom / Gardien / Numéro / Alias (optionnelle).
    Alias sert uniquement pour rattacher un buteur de la base buts quand son nom y est écrit
    différemment que dans la feuille de tag (ex. compo='DONGO', base buts='DONGMO')."""
    wb = openpyxl.load_workbook(path, data_only=True)
    compo = {}  # team -> list of dict(nom, numero, gardien, alias)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        team = str(ws.cell(row=1, column=2).value).strip()
        players = []
        header_row = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value == "Nom")
        header_cells = [str(c.value).strip() if c.value else "" for c in ws[header_row]]
        alias_col = next((i for i, h in enumerate(header_cells) if h.lower().startswith("alias")), None)
        for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row):
            nom = row[0].value
            if nom is None:
                continue
            gardien = str(row[1].value).strip().upper() == "OUI"
            numero_brut = row[2].value
            if numero_brut is None or str(numero_brut).strip() == "":
                raise CompoError(
                    f"Compo '{team}' : le joueur '{nom}' n'a pas de numéro renseigné (colonne Numéro vide, "
                    f"ligne {row[0].row}). Complète-la avant de régénérer le rapport."
                )
            try:
                numero = int(numero_brut)
            except (TypeError, ValueError):
                raise CompoError(
                    f"Compo '{team}' : le numéro du joueur '{nom}' n'est pas un nombre valide "
                    f"('{numero_brut}', ligne {row[0].row})."
                )
            alias = str(row[alias_col].value).strip().upper() if alias_col is not None and row[alias_col].value else None
            players.append({"nom": str(nom).strip().upper(), "numero": numero, "gardien": gardien, "alias": alias})
        compo[team] = players
    return compo


def load_raw(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Sheet1"]
    header = [c.value for c in ws[1]]
    rows = []
    for r in range(2, ws.max_row + 1):
        label = ws.cell(row=r, column=1).value
        if label is None:
            continue
        values = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        rows.append((label, values))
    return header, rows


def build_column_map(header, compo):
    """col_index (0-based dans la liste header, où 0=label) -> (team, player_key) ou type special."""
    col_map = {}
    special_cols = {}
    all_players = {}  # (numero, nom) -> team
    for team, players in compo.items():
        for p in players:
            all_players[(p["numero"], p["nom"])] = team

    header_re = re.compile(r"^\s*(\d+)\.\s*(.+?)\s*$")
    unmatched = []
    for idx, h in enumerate(header):
        if idx == 0 or h is None:
            continue
        if h in ("Gardien Volant", "Power Play"):
            special_cols[idx] = h
            continue
        if h in ("But LAVAL", "But NICE") or h.startswith("But "):
            special_cols[idx] = h  # qualificatif "action ayant mené à un but", hors scope v1
            continue
        m = header_re.match(str(h))
        if not m:
            unmatched.append((idx, h))
            continue
        numero = int(m.group(1))
        nom = m.group(2).strip().upper()
        team = all_players.get((numero, nom))
        if team is None:
            unmatched.append((idx, h))
            continue
        col_map[idx] = (team, nom)
    return col_map, special_cols, unmatched


class CompoError(Exception):
    """Levée quand la compo et la feuille brute ne matchent pas parfaitement."""
    pass


def validate_compo_vs_raw(header, compo):
    """
    Contrôle un seul sens : une colonne joueur de la feuille brute (donc quelqu'un qui a vraiment
    joué/tagué) doit être reconnue dans la compo, sinon ses stats disparaîtraient silencieusement
    des totaux (cas DUQUE). L'inverse (un joueur de la compo absent de la feuille brute) est normal :
    la compo est l'effectif du club, pas la feuille de match — un remplaçant qui n'a pas joué n'a
    simplement aucune colonne, ce n'est pas une erreur.
    """
    header_re = re.compile(r"^\s*(\d+)\.\s*(.+?)\s*$")
    raw_players = {}  # (numero, nom) -> libellé brut d'origine
    for h in header:
        if h is None or h in ("Gardien Volant", "Power Play") or str(h).startswith("But "):
            continue
        m = header_re.match(str(h))
        if m:
            raw_players[(int(m.group(1)), m.group(2).strip().upper())] = h

    compo_keys = {(p["numero"], p["nom"]): (team, p["nom"]) for team, players in compo.items() for p in players}
    raw_absents = [(key, label) for key, label in raw_players.items() if key not in compo_keys]

    if raw_absents:
        msg = ["Joueur(s) tagué(s) dans la feuille brute mais absent(s) de la compo :"]
        for (numero, nom), label in raw_absents:
            msg.append(f"  - La feuille brute contient la colonne '{label}' mais ce joueur"
                       f" n'est dans aucune des feuilles de la compo (joueur à ajouter).")
        raise CompoError("\n".join(msg))
    return True


class GoalsDbError(Exception):
    """Levée quand un but de la base buteurs ne peut pas être rattaché à un joueur de la compo."""
    pass


# Mapping nom court (compo / feuille de tag) -> nom officiel complet (base buteurs)
TEAM_NAME_MAP = {
    "GOAL": "GOAL FUTSAL CLUB", "NANTES": "NANTES METROPOLE F.", "ACASA": "PARIS ACASA",
    "LAVAL": "ETOILE LAVALLOISE FC", "MTP": "MONTPELLIER MED. F.", "SPORTING": "SPORTING CLUB PARIS",
    "AVION": "AS AVION FUTSAL", "KINGERSHEIM": "FC KINGERSHEIM", "NICE": "NICE FUTSAL CLUB",
    "TOULON": "TOULON METROPOLE F.", "TOULOUSE": "UJS TOULOUSE",
}


def match_scorer_name(full_name, compo_players):
    """Fait correspondre un nom complet ('SOUFIANE EL MESRAR') à un joueur de la compo :
    tous les mots du nom compo, OU de son alias, doivent apparaître dans le nom complet.
    compo_players : liste de dicts {nom, alias}. Retourne le nom canonique (compo) ou None si 0/2+ matchs."""
    full_words = set(full_name.upper().split())
    matches = []
    for p in compo_players:
        candidates = [p["nom"]] + ([p["alias"]] if p.get("alias") else [])
        if any(set(c.upper().split()).issubset(full_words) for c in candidates):
            matches.append(p["nom"])
    return matches[0] if len(matches) == 1 else None


def _origine_from_team_sheet(wb, team_short, journee, minute, scorer_full):
    """Repli : si l'origine est vide dans l'onglet maître, on va la chercher dans l'onglet de
    l'équipe qui a marqué (parfois renseignée là seulement, comme observé sur la saison en cours)."""
    if team_short not in wb.sheetnames:
        return None
    ws = wb[team_short]
    header = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(header) if h}
    needed = {"journee", "minute", "joueur", "origine_but"}
    if not needed.issubset(col):
        return None
    for r in range(2, ws.max_row + 1):
        row = [ws.cell(row=r, column=c + 1).value for c in range(len(header))]
        if row[col["journee"]] == journee and row[col["minute"]] == minute and row[col["joueur"]] == scorer_full:
            return (row[col["origine_but"]] or "").strip()
    return None


def get_home_away(goals_db_path, team_a, team_b, journee):
    """Détermine qui est domicile/extérieur pour CE match via l'onglet maître de la base buteurs
    (colonnes equipe_domicile/equipe_exterieure). Repli sur (team_a, team_b) tel quel si le match
    n'y est pas trouvé (base pas encore à jour, etc.)."""
    off_a, off_b = TEAM_NAME_MAP.get(team_a), TEAM_NAME_MAP.get(team_b)
    if off_a is None or off_b is None:
        return team_a, team_b
    wb = openpyxl.load_workbook(goals_db_path, data_only=True)
    sheet_name = next((s for s in wb.sheetnames if s.strip().upper().startswith("BUT D1")), None)
    if sheet_name is None:
        return team_a, team_b
    ws = wb[sheet_name]
    header = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(header) if h}
    if not {"journee", "equipe_domicile", "equipe_exterieure"}.issubset(col):
        return team_a, team_b
    for r in range(2, ws.max_row + 1):
        dom = ws.cell(row=r, column=col["equipe_domicile"] + 1).value
        ext = ws.cell(row=r, column=col["equipe_exterieure"] + 1).value
        jr = ws.cell(row=r, column=col["journee"] + 1).value
        if jr == journee and {dom, ext} == {off_a, off_b}:
            return (team_a, team_b) if dom == off_a else (team_b, team_a)
    return team_a, team_b


def load_goals_for_match(goals_db_path, team_a, team_b, journee, compo):
    """Filtre l'onglet maître 'BUT D1' par journée + les 2 équipes (noms courts de la compo),
    rattache chaque buteur à un nom de la compo, et retourne (buteurs_a, buteurs_b) triés par minute."""
    off_a, off_b = TEAM_NAME_MAP.get(team_a), TEAM_NAME_MAP.get(team_b)
    if off_a is None or off_b is None:
        raise GoalsDbError(f"Nom officiel introuvable pour '{team_a}' ou '{team_b}' — ajoute-le à TEAM_NAME_MAP.")

    wb = openpyxl.load_workbook(goals_db_path, data_only=True)
    sheet_name = next((s for s in wb.sheetnames if s.strip().upper().startswith("BUT D1")), None)
    if sheet_name is None:
        raise GoalsDbError("Onglet maître 'BUT D1' introuvable dans le fichier buteurs.")
    ws = wb[sheet_name]
    header = [c.value for c in ws[1]]
    col = {h: i for i, h in enumerate(header) if h}

    players_a = [{"nom": p["nom"], "alias": p.get("alias")} for p in compo[team_a]]
    players_b = [{"nom": p["nom"], "alias": p.get("alias")} for p in compo[team_b]]

    buteurs_a, buteurs_b, unmatched = [], [], []
    for r in range(2, ws.max_row + 1):
        row = [ws.cell(row=r, column=c + 1).value for c in range(len(header))]
        if row[col["journee"]] != journee:
            continue
        dom, ext = row[col["equipe_domicile"]], row[col["equipe_exterieure"]]
        if {dom, ext} != {off_a, off_b}:
            continue
        marque = row[col["equipe_marque"]]
        scorer_full = row[col["joueur"]]
        minute = row[col["minute"]]
        origine = ((row[col["origine_but"]] or "") if "origine_but" in col else "").strip()
        team_short = team_a if marque == off_a else team_b
        if not origine:
            origine = _origine_from_team_sheet(wb, team_short, journee, minute, scorer_full) or ""
        candidates = players_a if team_short == team_a else players_b
        matched = match_scorer_name(scorer_full, candidates)
        entry = {"minute": f"{minute}'", "joueur": matched or f"?? ({scorer_full})", "origine": origine}
        if matched is None:
            unmatched.append(scorer_full)
        (buteurs_a if team_short == team_a else buteurs_b).append(entry)

    if not buteurs_a and not buteurs_b:
        raise GoalsDbError(f"Aucun but trouvé pour {team_a} vs {team_b} à la journée {journee} dans la base buteurs.")
    if unmatched:
        raise GoalsDbError(
            "Buteur(s) non reconnu(s) dans la compo (nom à corriger dans l'un des deux fichiers) : "
            + ", ".join(unmatched)
        )
    buteurs_a.sort(key=lambda g: int(g["minute"].rstrip("'")))
    buteurs_b.sort(key=lambda g: int(g["minute"].rstrip("'")))
    return buteurs_a, buteurs_b


def build_row_index(rows, teams):
    """(categorie_normalisee, team_ou_None) -> (values_list, total)"""
    idx = {}
    for label, values in rows:
        label_clean = " ".join(str(label).split())  # normalise espaces multiples
        matched_team = None
        for team in teams:
            if label_clean.upper().endswith(" " + team.upper()):
                matched_team = team
                category = label_clean[: -(len(team) + 1)].strip()
                break
        if matched_team is None:
            category = label_clean
        idx[(category, matched_team)] = values
    return idx


def get_val(row_index, category, team, col_idx):
    values = row_index.get((category, team))
    if values is None:
        return 0
    v = values[col_idx]
    return v if isinstance(v, (int, float)) else 0


def get_total(row_index, category, team):
    values = row_index.get((category, team))
    if values is None:
        return 0
    v = values[-1]
    return v if isinstance(v, (int, float)) else 0


def detect_teams_in_raw(rows, known_teams):
    """Déduit les équipes réellement présentes dans CE match à partir des suffixes des libellés de ligne,
    pour ignorer les autres équipes de la compo (qui grossit au fil des matchs, style carnet de rosters)."""
    present = set()
    for label, _values in rows:
        label_clean = " ".join(str(label).split()).upper()
        for team in known_teams:
            if label_clean.endswith(" " + team.upper()):
                present.add(team)
    if len(present) != 2:
        raise CompoError(
            f"Impossible de déduire les 2 équipes du match depuis la feuille brute (trouvé : {sorted(present)}). "
            f"Vérifie que le nom d'équipe dans la compo correspond exactement au nom utilisé dans les libellés "
            f"de la feuille de tag (ex. 'LAVAL' partout, pas 'Laval' ou 'FC Laval')."
        )
    return present


def compute(raw_path=RAW_PATH, compo_path=COMPO_PATH):
    compo_full = load_compo(compo_path)
    header, rows = load_raw(raw_path)
    present = detect_teams_in_raw(rows, list(compo_full.keys()))
    # ordre déterministe : celui des feuilles de la compo, pas l'ordre (aléatoire) du set détecté
    teams = [t for t in compo_full.keys() if t in present]
    compo = {team: compo_full[team] for team in teams}
    validate_compo_vs_raw(header, compo)  # lève CompoError et stoppe tout si incohérence
    col_map, special_cols, unmatched = build_column_map(header, compo)
    row_index = build_row_index(rows, teams)

    # stats indiv
    player_stats = {}
    for idx, (team, nom) in col_map.items():
        opp = [t for t in teams if t != team][0]
        gardien = any(p["gardien"] for p in compo[team] if p["nom"] == nom)

        tir_cadre = get_val(row_index, "Tir Cadré", team, idx)
        tir_hc = get_val(row_index, "Tir HC", team, idx)
        tir_contre = get_val(row_index, "Tir Contré", team, idx)
        buts = get_val(row_index, "But", team, idx)

        duel_off_g = get_val(row_index, "Duel Gagné OFF", team, idx)
        duel_off_p = get_val(row_index, "Duel Perdu OFF", team, idx)
        duel_def_p = get_val(row_index, "Duel Gagné OFF", opp, idx)   # adversaire gagne OFF -> je perds DEF
        duel_def_g = get_val(row_index, "Duel Perdu OFF", opp, idx)   # adversaire perd OFF -> je gagne DEF

        perte_solo = get_val(row_index, "Perte de balle", team, idx)
        ballon_rendu = get_val(row_index, "Ballon Rendu", team, idx)
        recup_adverse = get_val(row_index, "Récupération", opp, idx)
        perte_balle = perte_solo + ballon_rendu + recup_adverse

        recuperation = get_val(row_index, "Récupération", team, idx)
        interception = get_val(row_index, "Interception", team, idx)

        passe_loupee_solo = get_val(row_index, "Passe Loupée", team, idx)
        interception_adverse = get_val(row_index, "Interception", opp, idx)
        passe_loupee = passe_loupee_solo + interception_adverse

        passe_decisive = get_val(row_index, "Passe Désicive", team, idx)
        poteau = get_val(row_index, "Poteau", team, idx)

        faute_subie = get_val(row_index, "Faute Subie", team, idx)
        faute_commise = get_val(row_index, "Faute Subie", opp, idx)

        stats = {
            "equipe": team, "gardien": gardien,
            "tirs": tir_cadre + tir_hc + tir_contre,  # but est déjà inclus dans tir_cadre
            "tirs_cadres": tir_cadre, "tirs_hc": tir_hc, "tirs_contres": tir_contre, "buts": buts,
            "duel_off_gagne": duel_off_g, "duel_off_perdu": duel_off_p,
            "duel_def_gagne": duel_def_g, "duel_def_perdu": duel_def_p,
            "perte_balle": perte_balle, "recuperation": recuperation, "interception": interception,
            "passe_loupee": passe_loupee, "passe_decisive": passe_decisive, "poteau": poteau,
            "faute_subie": faute_subie, "faute_commise": faute_commise,
        }

        if gardien:
            # "subis" = valeurs cross-taguées sur la ligne de l'ÉQUIPE ADVERSE (celle qui tire/marque),
            # pas la ligne de sa propre équipe (qui donnerait ses propres tirs/buts marqués).
            tir_cadre_subi = get_val(row_index, "Tir Cadré", opp, idx)
            but_subi = get_val(row_index, "But", opp, idx)
            stats["arrets"] = tir_cadre_subi - but_subi
            stats["tirs_cadres_subis"] = tir_cadre_subi
            stats["buts_subis"] = but_subi
            stats["relance_facile"] = get_val(row_index, "Relance Facile Réussi", team, idx)
            stats["relance_diff_ok"] = get_val(row_index, "Relance Difficile Réussi", team, idx)
            stats["relance_diff_ko"] = get_val(row_index, "Relance Difficile Loupée", team, idx)

        player_stats[nom] = stats

    # stats équipe (volume only + somme des indiv)
    team_stats = {}
    for team in teams:
        opp = [t for t in teams if t != team][0]
        att_placee = (get_total(row_index, "ATT 4-0", team) + get_total(row_index, "ATT 1-2-1", team)
                      + get_total(row_index, "ATT", team))
        team_stats[team] = {
            "attaque_placee": att_placee,
            "defense": (get_total(row_index, "ATT 4-0", opp) + get_total(row_index, "ATT 1-2-1", opp)
                        + get_total(row_index, "ATT", opp)),
            "transition_off": get_total(row_index, "Trans OFF", team),
            "transition_def": get_total(row_index, "Trans OFF", opp),
            "touche_off": get_total(row_index, "Touche OFF", team),
            "touche_def": get_total(row_index, "Touche DEF", team),
            "corner": get_total(row_index, "Corner", team),
            "penalty": get_total(row_index, "Penalty", team),
            "jet_franc": get_total(row_index, "Jet Franc", team),
            "coup_franc": get_total(row_index, "Coup Franc", team),
        }

    # agrégation indiv -> équipe pour les stats sommables
    agg_fields = ["tirs", "tirs_cadres", "tirs_hc", "tirs_contres", "buts",
                  "duel_off_gagne", "duel_off_perdu", "duel_def_gagne", "duel_def_perdu",
                  "perte_balle", "recuperation", "interception", "passe_loupee",
                  "passe_decisive", "poteau", "faute_subie", "faute_commise"]
    for team in teams:
        for f in agg_fields:
            team_stats[team][f] = sum(s[f] for s in player_stats.values() if s["equipe"] == team)

    # volume gardien volant / power play : toutes lignes confondues, pas par équipe
    gv_total = sum(v for idx_c, h in special_cols.items() if h == "Gardien Volant"
                   for label, values in rows for v in [values[idx_c]] if isinstance(v, (int, float)))
    pp_total = sum(v for idx_c, h in special_cols.items() if h == "Power Play"
                   for label, values in rows for v in [values[idx_c]] if isinstance(v, (int, float)))

    return player_stats, team_stats, gv_total, pp_total


if __name__ == "__main__":
    player_stats, team_stats, gv_total, pp_total = compute()

    print("\n=== STATS ÉQUIPE ===")
    fields = ["buts", "tirs", "tirs_cadres", "tirs_hc", "tirs_contres",
              "duel_off_gagne", "duel_off_perdu", "duel_def_gagne", "duel_def_perdu",
              "perte_balle", "recuperation", "interception", "faute_subie", "faute_commise",
              "attaque_placee", "defense", "transition_off", "transition_def",
              "touche_off", "touche_def", "corner", "coup_franc"]
    header_line = "STAT".ljust(18) + "".join(t.rjust(10) for t in team_stats.keys())
    print(header_line)
    for f in fields:
        print(f.ljust(18) + "".join(str(team_stats[t].get(f, 0)).rjust(10) for t in team_stats.keys()))
    print(f"Gardien Volant (volume total match): {gv_total}")
    print(f"Power Play (volume total match): {pp_total}")

    print("\n=== STATS INDIV (joueurs de champ) ===")
    cols = ["equipe", "tirs", "tirs_cadres", "buts", "duel_off_gagne", "duel_off_perdu",
            "duel_def_gagne", "duel_def_perdu", "perte_balle", "recuperation", "interception",
            "faute_subie", "faute_commise"]
    print("JOUEUR".ljust(16) + "".join(c[:10].rjust(9) for c in cols))
    for nom, s in player_stats.items():
        if s["gardien"]:
            continue
        print(nom.ljust(16) + "".join(str(s.get(c, 0)).rjust(9) for c in cols))

    print("\n=== GARDIENS ===")
    for nom, s in player_stats.items():
        if s["gardien"]:
            print(f"{nom} ({s['equipe']}): {s['arrets']} arrêts / {s['tirs_cadres_subis']} tirs cadrés subis, "
                  f"{s['buts_subis']} buts subis, relances {s['relance_facile']}F/{s['relance_diff_ok']}D+/{s['relance_diff_ko']}D-")
